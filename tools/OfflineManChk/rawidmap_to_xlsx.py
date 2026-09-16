#!/usr/bin/env python3
"""Convert a *_rawidmap.txt into an Excel (.xlsx) table.

Usage:
    python rawidmap_to_xlsx.py <btrklist<n>_rawidmap.txt> [output.xlsx]

<*_rawidmap.txt> is one line per track:
    PL eventid rawid1 zone rawid2 ax ay x y ph1 ph2 trk_type
(the format RawidMapChecker.py reads/writes). trk_type is 0 (real track),
1 (extrapolated), or -1 (reference track).

Writes an .xlsx workbook with three sheets:
    "実在_外挿" -> data rows with trk_type 0 or 1
    "Reference" -> data rows with trk_type -1
    "Classification" -> the fixed list of classification choices (below)
Each data sheet's header row (row 1) is:
    PL, eventid, trktype, ax, ay, x, y, ph1, ph2, Classification
followed by one row per matching line in the input file, in file order
(duplicates in the input are kept as separate rows). The "Classification"
column is left blank for manual annotation, with a dropdown (Excel data
validation) restricting it to the choices listed on the "Classification"
sheet.
"""
import argparse
import os

from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation

HEADER = ["PL", "eventid", "trktype", "ax", "ay", "x", "y", "ph1", "ph2", "Classification"]

CLASSIFICATION_CHOICES = [
    "Both exist.",
    "Only the lens side exists.",
    "Only the stage side exists.",
    "Neither exists.",
    "Cannot judge",
    "?",
]


def read_rows(rawidmap_path):
    real_pred_rows = []
    ref_rows = []
    with open(rawidmap_path, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            parts = line.split()
            if len(parts) != 12:
                print(f"Skipping malformed line: {line!r}")
                continue
            pl, eventid = int(parts[0]), int(parts[1])
            ax, ay, x, y, ph1, ph2 = (float(v) for v in parts[5:11])
            trk_type = int(parts[11])
            row = [pl, eventid, trk_type, ax, ay, x, y, ph1, ph2, None]
            (ref_rows if trk_type == -1 else real_pred_rows).append(row)
    return real_pred_rows, ref_rows


def write_sheet(wb, title, rows):
    ws = wb.create_sheet(title)
    ws.append(HEADER)
    for row in rows:
        ws.append(row)

    if rows:
        dv = DataValidation(type="list", formula1="Classification!$B$2:$B$7", allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"J2:J{len(rows) + 1}")
    return ws


def write_classification_sheet(wb):
    ws = wb.create_sheet("Classification")
    ws.append([None, "manualcheck results"])
    for i, choice in enumerate(CLASSIFICATION_CHOICES, start=1):
        ws.append([i, choice])
    return ws


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("rawidmap_txt", help=r"e.g. ...\btrklist6_rawidmap.txt")
    parser.add_argument("output_xlsx", nargs="?", help="default: <rawidmap stem>.xlsx next to the input file")
    args = parser.parse_args()

    rawidmap_path = os.path.abspath(args.rawidmap_txt)
    if args.output_xlsx:
        output_path = os.path.abspath(args.output_xlsx)
    else:
        stem = os.path.splitext(os.path.basename(rawidmap_path))[0]
        output_path = os.path.join(os.path.dirname(rawidmap_path), f"{stem}.xlsx")

    real_pred_rows, ref_rows = read_rows(rawidmap_path)

    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "実在_外挿", real_pred_rows)
    write_sheet(wb, "Reference", ref_rows)
    write_classification_sheet(wb)
    wb.save(output_path)

    print(f"input : {rawidmap_path}")
    print(f"output: {output_path}")
    print(f"実在_外挿: {len(real_pred_rows)} 行")
    print(f"Reference: {len(ref_rows)} 行")


if __name__ == "__main__":
    main()
